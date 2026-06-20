"""Deterministic safe fixtures for every Phase 11 parser."""

from __future__ import annotations

from io import BytesIO

import pytest
from app.services.documents.parsers.language import detect_language
from app.services.documents.parsers.registry import DocumentParserRegistry
from app.services.documents.parsers.types import ParseFailure, SafeParseErrorCode
from docx import Document as DocxDocument
from openpyxl import Workbook  # type: ignore[import-untyped]


@pytest.fixture
def registry() -> DocumentParserRegistry:
    return DocumentParserRegistry(maximum_characters=10_000, maximum_sections=100)


def test_pdf_parser_preserves_page_order_and_page_count(registry: DocumentParserRegistry) -> None:
    parsed = registry.parse(file_type="pdf", payload=_pdf_with_text("Hello PDF"))

    assert "Hello PDF" in parsed.extracted_text
    assert parsed.page_count == 1
    assert parsed.spans[0].page_number == 1
    assert parsed.spans[0].kind == "page"


def test_docx_parser_associates_body_with_heading(registry: DocumentParserRegistry) -> None:
    document = DocxDocument()
    document.add_heading("Summary", level=1)
    document.add_paragraph("Safe document body.")
    buffer = BytesIO()
    document.save(buffer)

    parsed = registry.parse(file_type="docx", payload=buffer.getvalue())

    assert parsed.extracted_text == "Safe document body."
    assert parsed.spans[0].label == "Summary"


@pytest.mark.parametrize(
    ("file_type", "payload", "expected_label"),
    [
        ("txt", b"One\r\nTwo\r\n", None),
        ("markdown", b"# First\nAlpha\n## Second\nBeta", "First"),
        ("csv", b"name,amount\nAda,12", "Header"),
        (
            "eml",
            b"From: sender@example.invalid\nSubject: Synthetic\n\nPlain body wins.",
            "Headers",
        ),
    ],
)
def test_textual_parsers_produce_normalized_location_metadata(
    registry: DocumentParserRegistry,
    file_type: str,
    payload: bytes,
    expected_label: str | None,
) -> None:
    parsed = registry.parse(file_type=file_type, payload=payload)

    assert parsed.extracted_text
    assert parsed.spans[0].start == 0
    assert parsed.spans[0].end <= len(parsed.extracted_text)
    assert parsed.spans[0].label == expected_label


def test_xlsx_parser_preserves_worksheet_and_saved_formula(
    registry: DocumentParserRegistry,
) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Budget"
    worksheet.append(["Amount", "Tax"])
    worksheet.append([10, "=A2*0.25"])
    buffer = BytesIO()
    workbook.save(buffer)

    parsed = registry.parse(file_type="xlsx", payload=buffer.getvalue())

    assert "=A2*0.25" in parsed.extracted_text
    assert parsed.spans[0].label == "Budget"
    assert parsed.spans[0].kind == "worksheet"


def test_eml_parser_uses_html_only_when_plain_text_is_unavailable(
    registry: DocumentParserRegistry,
) -> None:
    payload = (
        b"From: sender@example.invalid\nSubject: Synthetic\nContent-Type: text/html\n\n"
        b"<p>Fallback <strong>body</strong></p>"
    )

    parsed = registry.parse(file_type="eml", payload=payload)

    assert "Fallback body" in parsed.extracted_text


@pytest.mark.parametrize(
    ("file_type", "payload", "code"),
    [
        ("unknown", b"safe", SafeParseErrorCode.UNSUPPORTED_TYPE),
        ("txt", b"\n\r", SafeParseErrorCode.EMPTY_TEXT),
        ("pdf", b"not a pdf", SafeParseErrorCode.MALFORMED_DOCUMENT),
    ],
)
def test_malformed_or_unreadable_inputs_map_to_safe_codes(
    registry: DocumentParserRegistry,
    file_type: str,
    payload: bytes,
    code: SafeParseErrorCode,
) -> None:
    with pytest.raises(ParseFailure) as raised:
        registry.parse(file_type=file_type, payload=payload)

    assert raised.value.code == code
    assert payload.decode("utf-8", errors="ignore") not in raised.value.summary


def test_parser_enforces_text_and_section_bounds() -> None:
    text_limited = DocumentParserRegistry(maximum_characters=4, maximum_sections=10)
    section_limited = DocumentParserRegistry(maximum_characters=100, maximum_sections=1)

    with pytest.raises(ParseFailure, match="text_limit_exceeded"):
        text_limited.parse(file_type="txt", payload=b"sixxx")
    with pytest.raises(ParseFailure, match="section_limit_exceeded"):
        section_limited.parse(file_type="markdown", payload=b"# One\na\n# Two\nb")


def test_language_detection_is_deterministic_for_english_norwegian_and_short_text() -> None:
    english = detect_language(
        "This is a longer English sentence used only for deterministic parser language testing. "
        * 3,
        minimum_characters=40,
        confidence_threshold=0.8,
    )
    norwegian = detect_language(
        "Dette er en lengre norsk bokmaltekst som brukes for sikker sprakkjenkjenning i testen. "
        * 3,
        minimum_characters=40,
        confidence_threshold=0.8,
    )
    short = detect_language("kort", minimum_characters=40, confidence_threshold=0.8)

    assert english.language == "en"
    assert norwegian.language == "nb"
    assert short.language == "unknown"


def _pdf_with_text(value: str) -> bytes:
    """Build a tiny one-page PDF without adding a test-only rendering dependency."""

    content = f"BT\n/F1 12 Tf\n72 720 Td\n({value}) Tj\nET".encode("ascii")
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        (
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            b"/Resources << /Font << /F1 5 0 R >> >> /Contents 4 0 R >>"
        ),
        b"<< /Length " + str(len(content)).encode() + b" >>\nstream\n" + content + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    output = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for index, object_value in enumerate(objects, 1):
        offsets.append(len(output))
        output.extend(f"{index} 0 obj\n".encode())
        output.extend(object_value)
        output.extend(b"\nendobj\n")
    xref_start = len(output)
    output.extend(f"xref\n0 {len(objects) + 1}\n".encode())
    output.extend(b"0000000000 65535 f \n")
    for offset in offsets[1:]:
        output.extend(f"{offset:010d} 00000 n \n".encode())
    output.extend(
        b"trailer\n<< /Size "
        + str(len(objects) + 1).encode()
        + b" /Root 1 0 R >>\nstartxref\n"
        + str(xref_start).encode()
        + b"\n%%EOF\n"
    )
    return bytes(output)
